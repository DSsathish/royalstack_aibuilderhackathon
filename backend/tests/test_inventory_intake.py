import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.connectors import csv_connector as csvc
from app.connectors import schema
from app.domain.models import Store
from app.main import app
from app.services import catalog, db

client = TestClient(app)
OWNER = {"X-Test-Uid": "owner-1"}


@pytest.fixture
def shop(store_id):
    db.store_ref(store_id).set(Store(store_id=store_id, name="S").model_dump())
    db.users().document("owner-1").set({"stores": [store_id], "roles": {store_id: "owner"}})
    catalog.invalidate(store_id)
    return store_id


def _xlsx(rows, headers=None):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(headers or schema.headers())
    for r in rows: ws.append(r)
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def _row(sku, name, price, on_hand, **extra):
    base = {"sku": sku, "name": name, "price": price, "on_hand": on_hand}
    return [base.get(h, extra.get(h, "")) for h in schema.headers()]


def test_schema_is_the_importer_contract():
    names = [f["name"] for f in schema.fields()]
    assert names == list(csvc.REQUIRED + csvc.OPTIONAL)
    assert [f["name"] for f in schema.fields() if f["required"]] == list(csvc.REQUIRED)
    assert schema.kind_of("price") == "number" and schema.kind_of("on_hand") == "integer" and schema.kind_of("brand") == "text"
    assert next(f for f in schema.fields() if f["name"] == "status")["choices"] == list(schema.STATUSES)
    check = schema.check_headers(["SKU", "Name", "Price", "On Hand", "Colour"])
    assert check["mapping"] == {"sku": "SKU", "name": "Name", "price": "Price", "on_hand": "On Hand"} and check["unknown"] == ["Colour"]
    assert schema.check_headers(["sku", "name"])["missing_required"] == ["price", "on_hand"]


def test_template_round_trips_through_upload(shop):
    res = client.get(f"/owner/{shop}/inventory/template.xlsx", headers=OWNER)
    assert res.status_code == 200 and res.headers["content-disposition"].endswith('"inventory-template.xlsx"')
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["Inventory"]
    assert [c.value for c in ws[1]] == schema.headers() and ws.max_row == 1 and "Guide" in wb.sheetnames
    ws.append(_row("ATTA-5KG", "Aashirvaad Atta 5 kg", 265, 40))
    ws.append(_row("MILK-500", "Aavin Toned Milk 500 ml", 28, 12))
    buf = io.BytesIO(); wb.save(buf)
    res = client.post(f"/owner/{shop}/inventory/upload", data={"apply": "false"}, files={"file": ("filled.xlsx", buf.getvalue())}, headers=OWNER)
    body = res.json()
    assert res.status_code == 200 and body["ok"] and body["row_count"] == 2 and body["preview"]["adds"] == 2 and body["result"] is None
    assert catalog.get_by_sku(shop, "ATTA-5KG") is None                                    # a check writes nothing
    res = client.post(f"/owner/{shop}/inventory/upload", data={"apply": "true"}, files={"file": ("filled.xlsx", buf.getvalue())}, headers=OWNER)
    assert res.json()["result"] == {"created": 2, "updated": 0, "skipped": 0, "errors": 0, "duplicates": 0, "file": "filled.xlsx"}
    assert catalog.get_by_sku(shop, "MILK-500").name == "Aavin Toned Milk 500 ml"
    assert any(a.to_dict()["action"] == "inventory_upload" for a in db.audit(shop).stream())


def test_upload_reports_actionable_problems_and_never_overwrites_by_default(shop):
    missing = _xlsx([["A1", "Atta"]], headers=["sku", "name"])
    body = client.post(f"/owner/{shop}/inventory/upload", files={"file": ("x.xlsx", missing)}, headers=OWNER).json()
    assert body["ok"] is False and body["missing_required"] == ["price", "on_hand"]
    rows = [_row("A1", "Atta", "abc", 4), _row("A2", "", 10, 1), _row("A3", "Rice", 90, "many"), _row("B1", "Dal", 120, 7), _row("b1", "Dal again", 100, 1)]
    body = client.post(f"/owner/{shop}/inventory/upload", data={"apply": "true"}, files={"file": ("x.xlsx", _xlsx(rows))}, headers=OWNER).json()
    assert body["ok"] and body["error_count"] == 3 and [e["line"] for e in body["errors"]] == [2, 3, 4]
    assert body["duplicates"] == [{"line": 6, "sku": "b1", "first_line": 5}]                # the first row wins
    assert body["result"]["created"] == 1 and catalog.get_by_sku(shop, "B1").price == 120
    again = client.post(f"/owner/{shop}/inventory/upload", data={"apply": "true"}, files={"file": ("x.csv", b"sku,name,price,on_hand\nB1,Dal premium,150,9\n")}, headers=OWNER).json()
    assert again["result"]["skipped"] == 1 and again["result"]["updated"] == 0 and catalog.get_by_sku(shop, "B1").price == 120
    allowed = client.post(f"/owner/{shop}/inventory/upload", data={"apply": "true", "allow_updates": "true"}, files={"file": ("x.csv", b"sku,name,price,on_hand\nB1,Dal premium,150,9\n")}, headers=OWNER).json()
    assert allowed["result"]["updated"] == 1 and catalog.get_by_sku(shop, "B1").price == 150
    assert client.post(f"/owner/{shop}/inventory/upload", files={"file": ("x.pdf", b"%PDF")}, headers=OWNER).status_code == 400
    assert client.get(f"/owner/{shop}/inventory/schema", headers={"X-Test-Uid": "stranger"}).status_code == 403


def test_manual_add_needs_a_sku_or_a_barcode(shop):
    assert client.post(f"/owner/{shop}/products", json={"name": "Loose Rice", "price": 60, "on_hand": 5}, headers=OWNER).status_code == 422
    res = client.post(f"/owner/{shop}/products", json={"sku": "RICE-LOOSE", "name": "Loose Rice", "price": 60, "on_hand": 5, "unit": "kg", "reorder_point": 2, "reorder_qty": 10, "subcategory": "Rice"}, headers=OWNER)
    assert res.status_code == 200, res.text
    view = catalog.get_by_sku(shop, "RICE-LOOSE")
    assert view.unit == "kg" and view.subcategory == "Rice"
