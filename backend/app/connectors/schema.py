"""The inventory fields, once. The Excel template, the Google Sheet layout, the upload check and the manual form all
read this table, which is built from the importer's own contract so the two can never disagree."""

import io

from app.connectors import csv_connector as csvc
from app.connectors.normalizer import MEANING

STATUSES = ("active", "inactive", "discontinued")
SIDES = ("left", "right")
_INTEGER = {"on_hand", "bay", "shelf", "reorder_point", "reorder_qty", "lead_time_days", "shelf_life_days"}
_NUMBER = {"price", "mrp", "cost_price", "gst_rate"}
_MEANING = {**MEANING, "pack_size": "pack size as sold (5 kg, 500 ml)", "cost_price": "what the shop pays per unit",
            "gst_rate": "GST rate in percent", "shelf_life_days": "shelf life in days", "status": "active, inactive or discontinued",
            "origin": "country or region of origin", "description": "longer description"}
_EXAMPLE = {"sku": "ATTA-5KG", "name": "Aashirvaad Whole Wheat Atta 5 kg", "price": "265", "on_hand": "40", "barcode": "8901030865278",
            "brand": "Aashirvaad", "category": "Grocery and Food Products", "subcategory": "Flours and Baking Ingredients",
            "description": "Whole wheat flour", "mrp": "280", "unit": "pc", "image_url": "", "tags": "atta|wheat flour",
            "substitutes": "ATTA-10KG", "aisle": "A1", "bay": "2", "shelf": "1", "side": "left", "reorder_point": "10",
            "reorder_qty": "30", "lead_time_days": "3", "supplier_id": "sup-itc", "pack_size": "5 kg", "cost_price": "240",
            "gst_rate": "5", "shelf_life_days": "180", "status": "active", "origin": "India"}


def kind_of(name: str) -> str:
    return "integer" if name in _INTEGER else "number" if name in _NUMBER else "text"


def fields() -> list[dict]:
    return [{"name": n, "required": n in csvc.REQUIRED, "kind": kind_of(n), "description": _MEANING.get(n, n.replace("_", " ")),
             "example": _EXAMPLE.get(n, ""), "choices": list(STATUSES) if n == "status" else list(SIDES) if n == "side" else []}
            for n in csvc.REQUIRED + csvc.OPTIONAL]


def headers() -> list[str]:
    return list(csvc.REQUIRED + csvc.OPTIONAL)


def _key(text: str) -> str:
    return " ".join(str(text).lower().replace("_", " ").replace("-", " ").split())


def check_headers(columns: list[str]) -> dict:
    """Match a file's columns to the fields, forgiving case, spaces and underscores. Unknown columns are ignored."""
    by_key = {_key(n): n for n in headers()}
    mapping, unknown = {}, []
    for col in columns:
        field = by_key.get(_key(col))
        if field and field not in mapping:
            mapping[field] = col
        elif not field:
            unknown.append(col)
    missing = [n for n in csvc.REQUIRED if n not in mapping]
    return {"mapping": mapping, "missing_required": missing, "unknown": unknown}


def template_xlsx() -> bytes:
    """An Excel template owners fill in: exact headers on the Inventory sheet (required ones highlighted, each with a
    note), dropdowns for status and side, and a Guide sheet. No sample row, so nothing imports by accident."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"
    required_fill = PatternFill("solid", fgColor="FDE2E1")
    optional_fill = PatternFill("solid", fgColor="EEF2F5")
    for col, f in enumerate(fields(), start=1):
        cell = ws.cell(row=1, column=col, value=f["name"])
        cell.font = Font(bold=True)
        cell.fill = required_fill if f["required"] else optional_fill
        cell.alignment = Alignment(vertical="center")
        cell.comment = Comment(("Required. " if f["required"] else "Optional. ") + f["description"] + (f" Example: {f['example']}" if f["example"] else ""), "Royal Stack")
        ws.column_dimensions[get_column_letter(col)].width = max(12, min(34, len(f["name"]) + 4))
    ws.freeze_panes = "A2"
    for name, choices in (("status", STATUSES), ("side", SIDES)):
        col = headers().index(name) + 1
        dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(col)}2:{get_column_letter(col)}5000")
    guide = wb.create_sheet("Guide")
    guide.append(["Column", "Required", "Type", "Meaning", "Example"])
    for c in guide[1]:
        c.font = Font(bold=True)
    for f in fields():
        guide.append([f["name"], "yes" if f["required"] else "no", f["kind"], f["description"], f["example"]])
    for col, width in zip("ABCDE", (18, 10, 10, 60, 34)):
        guide.column_dimensions[col].width = width
    guide.append([])
    guide.append(["Fill the Inventory sheet, keep the column names as they are, save as .xlsx or .csv, and upload it from the Inventory tab."])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
